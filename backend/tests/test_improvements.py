import io
import pytest
from docx import Document as DocxDoc
from pypdf import PdfWriter

from app.documents.chunking import chunk_extracted_document
from app.documents.embeddings import MockEmbeddingProvider
from app.documents.extractors import (
    DocumentExtractionError,
    extract_docx,
    extract_markdown,
    extract_pdf,
)
from app.evaluation.evaluator import calculate_metrics, match_findings_to_expected, run_evaluation


# =====================================================================
# 1. DOCUMENT EXTRACTION & RETRIEVAL TESTS
# =====================================================================

def test_markdown_extraction_and_chunking():
    md_content = """# Brand Guidelines
## Voice and Tone
Always maintain an editorial and confident voice. Avoid hype and superlatives.

## Terminology
Refer to users as customers. Do not use consumer or end-user.
""".encode("utf-8")

    result = extract_markdown(md_content)
    assert result.total_characters > 0
    assert "Voice and Tone" in result.headings
    assert "Terminology" in result.headings

    chunks = chunk_extracted_document(result, chunk_size=100, chunk_overlap=20)
    assert len(chunks) >= 2
    assert chunks[0].heading_path is not None
    assert chunks[0].chunk_index == 0


def test_docx_extraction():
    doc = DocxDoc()
    doc.add_heading("Product Strategy", level=1)
    doc.add_paragraph("Our editorial atelier reconciles drafts against brand standards.")
    doc.add_heading("Claims Policy", level=2)
    doc.add_paragraph("Every claim must be verified against ground truth.")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    result = extract_docx(docx_bytes)
    assert result.total_characters > 0
    assert "Product Strategy" in result.headings
    assert "Claims Policy" in result.headings
    assert any("editorial atelier" in s.content for s in result.sections)


def test_pdf_extraction():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    # Blank PDF has no text -> should produce explicit error
    buf = io.BytesIO()
    writer.write(buf)
    pdf_bytes = buf.getvalue()

    with pytest.raises(DocumentExtractionError) as exc_info:
        extract_pdf(pdf_bytes)
    assert "no extractable text" in str(exc_info.value).lower()


def test_invalid_and_oversized_uploads(client, creator_auth):
    # Empty file
    r = client.post(
        "/api/documents/upload",
        files={"file": ("empty.md", b"", "text/markdown")},
        headers=creator_auth,
    )
    assert r.status_code == 400

    # Unsupported extension
    r = client.post(
        "/api/documents/upload",
        files={"file": ("script.sh", b"echo hello", "text/x-shellscript")},
        headers=creator_auth,
    )
    assert r.status_code == 400

    # Corrupted PDF
    r = client.post(
        "/api/documents/upload",
        files={"file": ("fake.pdf", b"not a pdf content", "application/pdf")},
        headers=creator_auth,
    )
    assert r.status_code == 201
    doc_id = r.json()["id"]
    # Document status should show error
    detail = client.get(f"/api/documents/{doc_id}", headers=creator_auth)
    assert detail.json()["status"] == "error"


def test_document_upload_and_passage_retrieval(client, creator_auth):
    md_content = """# Brand Guidelines
## Editorial Voice
BrandPilot requires clear, precise editorial voice. Never use hyperbolic guarantees or sensational claims.
Approved terms require using customer instead of consumer.
""".encode("utf-8")

    res = client.post(
        "/api/documents/upload",
        files={"file": ("guidelines.md", md_content, "text/markdown")},
        headers=creator_auth,
    )
    assert res.status_code == 201
    doc_data = res.json()
    assert doc_data["status"] == "indexed"
    assert doc_data["extracted_metadata"]["chunk_count"] >= 1

    # Check passages endpoint
    passages_res = client.get(f"/api/documents/{doc_data['id']}/passages", headers=creator_auth)
    assert passages_res.status_code == 200
    passages = passages_res.json()
    assert len(passages) >= 1
    assert "Brand Guidelines" in passages[0]["heading_path"]

    # Delete document
    del_res = client.delete(f"/api/documents/{doc_data['id']}", headers=creator_auth)
    assert del_res.status_code == 200


@pytest.mark.asyncio
async def test_mock_embedding_provider_dimension():
    provider = MockEmbeddingProvider(dimension=1536)
    assert provider.dimension == 1536
    embeddings = await provider.embed_texts(["hello world", "brand voice guidelines"])
    assert len(embeddings) == 2
    assert len(embeddings[0]) == 1536
    assert len(embeddings[1]) == 1536


# =====================================================================
# 2. AUTHENTICATION, ROLES, AND DRAFT WORKFLOW TESTS
# =====================================================================

def test_auth_registration_and_login(client):
    email = "unique_user@brandpilot.test"
    reg = client.post("/api/auth/register", json={
        "email": email,
        "password": "SecurePassword123!",
        "full_name": "Atelier Reviewer",
        "role": "reviewer",
    })
    assert reg.status_code == 201
    assert "access_token" in reg.json()

    # Login with valid password
    login_res = client.post("/api/auth/login", json={
        "email": email,
        "password": "SecurePassword123!",
    })
    assert login_res.status_code == 200
    assert login_res.json()["user"]["role"] == "reviewer"

    # Login with incorrect password
    bad_login = client.post("/api/auth/login", json={
        "email": email,
        "password": "WrongPassword",
    })
    assert bad_login.status_code == 401


def test_role_based_access(client, creator_auth):
    # Creator cannot access reviewer queue
    res = client.get("/api/drafts/queue", headers=creator_auth)
    assert res.status_code == 403


def test_draft_submission_and_reviewer_approval_workflow(client, creator_auth, reviewer_auth):
    # 1. Creator creates draft
    create_res = client.post("/api/drafts", json={
        "title": "Summer Campaign Announcement",
        "draft_text": "Consumers can enjoy our services across all regions.",
    }, headers=creator_auth)
    assert create_res.status_code == 201
    draft = create_res.json()
    draft_id = draft["id"]
    assert draft["current_status"] == "draft"
    assert draft["latest_revision"]["revision_number"] == 1

    # 2. Creator submits draft
    sub_res = client.post(f"/api/drafts/{draft_id}/submit", headers=creator_auth)
    assert sub_res.status_code == 200
    assert sub_res.json()["current_status"] == "submitted"

    # 3. Reviewer queue lists the submitted draft
    q_res = client.get("/api/drafts/queue", headers=reviewer_auth)
    assert q_res.status_code == 200
    assert any(d["id"] == draft_id for d in q_res.json())

    # 4. UNAUTHORIZED SELF-APPROVAL: Creator cannot approve their own draft
    self_approve_res = client.post(
        f"/api/drafts/{draft_id}/decide",
        json={"decision": "approved", "comment": "I approve my own draft"},
        headers=creator_auth,
    )
    assert self_approve_res.status_code == 403

    # 5. Reviewer approves draft
    approve_res = client.post(
        f"/api/drafts/{draft_id}/decide",
        json={"decision": "approved", "comment": "Meets brand voice standards."},
        headers=reviewer_auth,
    )
    assert approve_res.status_code == 200
    assert approve_res.json()["current_status"] == "approved"


def test_suggestion_actions_and_audit_trail(client, creator_auth):
    create_res = client.post("/api/drafts", json={
        "title": "Terminology Test Draft",
        "draft_text": "Every consumer must log in to their dashboard.",
    }, headers=creator_auth)
    draft_id = create_res.json()["id"]
    findings = create_res.json()["findings"]
    assert len(findings) > 0
    finding_id = findings[0]["id"]

    # Accept suggestion
    act_res = client.post(
        f"/api/drafts/{draft_id}/suggestions/{finding_id}",
        json={"action": "accept", "comment": "Changed consumer to customer"},
        headers=creator_auth,
    )
    assert act_res.status_code == 200
    assert act_res.json()["finding_status"] == "accepted"

    # Audit trail verification
    audit_res = client.get(f"/api/drafts/{draft_id}/audit", headers=creator_auth)
    assert audit_res.status_code == 200
    events = audit_res.json()
    assert any(e["action_type"] == "suggestion_accept" for e in events)


def test_draft_version_history_and_immutability(client, creator_auth):
    create_res = client.post("/api/drafts", json={
        "title": "Version History Draft",
        "draft_text": "First draft original copy.",
    }, headers=creator_auth)
    draft_id = create_res.json()["id"]

    # Edit draft -> creates revision 2
    update_res = client.put(f"/api/drafts/{draft_id}", json={
        "draft_text": "Second draft updated copy.",
    }, headers=creator_auth)
    assert update_res.status_code == 200
    assert update_res.json()["latest_revision"]["revision_number"] == 2

    # Check revision history list
    revs_res = client.get(f"/api/drafts/{draft_id}/revisions", headers=creator_auth)
    assert revs_res.status_code == 200
    revs = revs_res.json()
    assert len(revs) == 2
    assert revs[0]["draft_text"] == "First draft original copy."
    assert revs[1]["draft_text"] == "Second draft updated copy."

    # Restore revision 1 -> creates revision 3 with rev 1's content, rev 1 remains unchanged!
    restore_res = client.post(f"/api/drafts/{draft_id}/restore/{revs[0]['id']}", headers=creator_auth)
    assert restore_res.status_code == 200
    assert restore_res.json()["latest_revision"]["revision_number"] == 3
    assert restore_res.json()["latest_revision"]["draft_text"] == "First draft original copy."

    # Verify rev 1 and 2 remain unchanged
    revs_after = client.get(f"/api/drafts/{draft_id}/revisions", headers=creator_auth).json()
    assert len(revs_after) == 3
    assert revs_after[0]["draft_text"] == "First draft original copy."
    assert revs_after[1]["draft_text"] == "Second draft updated copy."


# =====================================================================
# 3. EVALUATION & ANALYTICS TESTS
# =====================================================================

def test_metrics_calculation_and_zero_denominators():
    # Normal case: TP=8, FP=2, FN=0 -> Precision=0.8, Recall=1.0, F1=0.8889
    m = calculate_metrics(8, 2, 0, 10)
    assert m.precision == 0.8
    assert m.recall == 1.0
    assert m.f1 == 0.8889

    # Zero denominator case: no predictions, no expected (perfect negative control)
    m_zero = calculate_metrics(0, 0, 0, 1)
    assert m_zero.precision == 1.0
    assert m_zero.recall == 1.0
    assert m_zero.f1 == 1.0

    # Missed all: TP=0, FP=0, FN=5
    m_miss = calculate_metrics(0, 0, 5, 5)
    assert m_miss.precision == 0.0
    assert m_miss.recall == 0.0
    assert m_miss.f1 == 0.0


def test_duplicate_prediction_handling():
    predictions = [
        {"category": "Terminology", "excerpt": "consumer"},
        {"category": "Terminology", "excerpt": "consumer"},  # Duplicate
    ]
    expected = [
        {"category": "Terminology", "term": "consumer"},
    ]
    tp, fp, fn = match_findings_to_expected(predictions, expected)
    # The duplicate prediction must count as a false positive, not a second true positive!
    assert tp == 1
    assert fp == 1
    assert fn == 0


@pytest.mark.asyncio
async def test_evaluation_pipeline_execution():
    report = await run_evaluation()
    assert report.aggregate.sample_count == 10
    assert report.aggregate.precision > 0.90
    assert report.aggregate.recall > 0.80
    assert report.aggregate.f1 > 0.85
    assert "Terminology" in report.by_category


def test_recurring_issues_analytics_endpoint(client, creator_auth):
    res = client.get("/api/analytics/recurring-issues", headers=creator_auth)
    assert res.status_code == 200
    data = res.json()
    assert "summary" in data
    assert "most_triggered_rules" in data
    assert "category_distribution" in data
    assert "activity_timeline" in data
